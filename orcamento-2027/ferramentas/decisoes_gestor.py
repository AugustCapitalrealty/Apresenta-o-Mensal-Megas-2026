"""Decisões do gestor nas planilhas de comparação 2026 × 2027 → 23_DecisoesGestor.gs. Criado em 07/10/2026.

O gestor marca SIM/NÃO (coluna K, COMPARA?) e comenta (coluna L) nas planilhas
"ORÇAMENTO 2027 - COMPARAÇÃO DE ITENS RITMO 2026 x ORÇ 2027 - MEGA <X>" (base ritmo, desde 07/10/2026: colunas
A–P, o orçado 2026 em M e N) ou, na versão anterior, "… COMPARAÇÃO DE ITENS 2026 x 2027 - MEGA <X>" (colunas A–L). O Apps Script não lê .xlsx e uma delas é Planilha
Google: este script lê as planilhas em .xlsx e grava o que o gerador usa num arquivo .gs do projeto. Depois:
teste, commit e `clasp push`.

Uso (de dentro de orcamento-2027; precisa de `python -m pip install openpyxl`):
  python ferramentas/decisoes_gestor.py <planilha.xlsx> [<planilha.xlsx> ...]
Sem argumentos, lê as .xlsx "… RITMO 2026 x ORÇ 2027 - MEGA <X>.xlsx" da pasta "01 - CONTROLE DA APRESENTAÇÃO" e, de
Downloads, a versão mais nova de cada Mega que tiver sido baixada de Planilha Google (Arquivo → Fazer download → .xlsx).

O que vai para o .gs, por Mega:
  adiados → obras de 2026 que o gestor disse que não foram feitas e passaram para 2027 (NÃO + comentário com
            "deslocamos" ou "realocado"): a descrição e o valor do Orç 2026 (base ritmo: colunas M e N), o ritmo 2026
            e os itens de 2027 (nomes do modelo 090).
  pares   → os pares que ele marcou SIM, com o comentário (para a comparação item a item).
  semPar  → (só na base ritmo, desde 09/10/2026) os gastos do ritmo 2026 que não são par SIM nem obra adiada: o que
            "não se repete em 2027" no slide das demais variações (V16, 24_PorQueSobe.gs).
"""
import glob, json, os, re, sys
from datetime import date
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
PROJETO = os.path.dirname(AQUI)
PASTA_CONTROLE = os.path.join(PROJETO, '..', '..', '01 - CONTROLE DA APRESENTAÇÃO')
DOWNLOADS = os.path.join(os.path.expanduser('~'), 'Downloads')
SAIDA = os.path.join(PROJETO, '23_DecisoesGestor.gs')
ADIADO = re.compile(r'desloc|realoc', re.I)


def mega_do_arquivo(caminho):
    m = re.search(r'MEGA ([A-ZÁÍÉÓÚÃÕÇ]+)(?: \(\d+\))?\.xlsx$', os.path.basename(caminho), re.I)
    if not m: raise SystemExit('Nome fora do padrão (… - MEGA <X>.xlsx): ' + caminho)
    return 'Mega ' + m.group(1).capitalize()


def itens(txt):
    return [] if not txt or str(txt).strip() in ('—', '-') else [p.strip() for p in str(txt).split(' + ') if p.strip()]


def ler(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True)['Comparação']
    ritmo = str(ws.cell(row=5, column=4).value or '').startswith('Ritmo')
    adiados, pares, sem_par = [], [], []
    for r in ws.iter_rows(min_row=6, values_only=True):
        n, tipo, it26, o26, it27, cat, o27, d, leit, pq, comp, com, itOrc, orc = (list(r) + [None] * 14)[:14]
        if not isinstance(n, (int, float)): continue
        comp = str(comp or '').strip().upper()
        com = str(com or '').strip()
        if ritmo:
            orcado = str(itOrc or '').strip()
            orcado = '' if orcado in ('—', '-') else orcado
            de2026 = orcado or ('' if str(it26 or '').strip() in ('—', '-') else str(it26 or '').strip())
            reg = {'linha': int(n), 'de2026': de2026, 'orc2026': round(float(orc or 0), 2),
                   'ritmo2026': round(float(o26 or 0), 2), 'itens2027': itens(it27), 'orc2027': round(float(o27 or 0), 2),
                   'comentario': com}
        else:
            reg = {'linha': int(n), 'de2026': str(it26 or '').strip(), 'orc2026': round(float(o26 or 0), 2),
                   'itens2027': itens(it27), 'orc2027': round(float(o27 or 0), 2), 'comentario': com}
        if comp in ('NÃO', 'NAO') and ADIADO.search(com): adiados.append(reg)
        elif comp == 'SIM': pares.append(reg)
        elif ritmo and float(o26 or 0) > 0.5:
            sem_par.append({'linha': int(n), 'de2026': str(it26 or '').strip(), 'ritmo2026': round(float(o26), 2)})
    return adiados, pares, sem_par


def main(arqs):
    sys.stdout.reconfigure(encoding='utf-8')
    if not arqs:
        nome = 'ORÇAMENTO 2027 - COMPARAÇÃO DE ITENS RITMO 2026 x ORÇ 2027 - MEGA '
        por_mega = {}
        for a in sorted(glob.glob(os.path.join(PASTA_CONTROLE, nome + '*.xlsx'))) + \
                 sorted(glob.glob(os.path.join(DOWNLOADS, nome + '*.xlsx')), key=os.path.getmtime):
            por_mega[mega_do_arquivo(a)] = a      # o de Downloads (mais novo) vence
        arqs = list(por_mega.values())
    dados = {}
    for a in arqs:
        mega = mega_do_arquivo(a)
        adiados, pares, sem_par = ler(a)
        dados[mega] = {'arquivo': os.path.basename(a), 'adiados': adiados, 'pares': pares, 'semPar': sem_par}
        print(f'{mega}: {len(adiados)} obras adiadas de 2026 (R$ {sum(x["orc2027"] for x in adiados):,.0f} no Orç 2027, '
              f'R$ {sum(x.get("ritmo2026", 0) for x in adiados):,.0f} no ritmo 2026), '
              f'{len(pares)} pares SIM, {len(sem_par)} gastos de 2026 sem par '
              f'(R$ {sum(x["ritmo2026"] for x in sem_par):,.0f})  ← {os.path.basename(a)}')
    corpo = json.dumps(dados, ensure_ascii=False, indent=1)
    with open(SAIDA, 'w', encoding='utf-8', newline='\n') as f:
        f.write('/**\n * ARQUIVO: 23_DecisoesGestor.gs — GERADO por ferramentas/decisoes_gestor.py em ' +
                date.today().strftime('%d/%m/%Y') + '. Não edite à mão:\n'
                ' * o gestor muda a planilha de comparação, rode o script de novo.\n'
                ' * Decisões do gestor nas planilhas "ORÇAMENTO 2027 - COMPARAÇÃO DE ITENS 2026 x 2027 - MEGA <X>",\n'
                ' * por Mega: obras de 2026 adiadas para 2027, os pares que ele marcou SIM e os gastos do ritmo 2026 sem par.\n */\n')
        f.write('const ORC_DECISOES_GESTOR = ' + corpo + ';\n')
    print('ok ->', os.path.relpath(SAIDA, PROJETO))


if __name__ == '__main__':
    main(sys.argv[1:])
