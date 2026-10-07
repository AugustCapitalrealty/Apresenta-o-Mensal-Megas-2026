import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

S = os.path.dirname(os.path.abspath(__file__))
# Pares revisados por curadoria.js (node ferramentas/curadoria.js . ferramentas/comparacao_linhas_curitiba.json)
linhas = json.load(open(os.path.join(S, 'comparacao_linhas_curitiba.json'), encoding='utf-8'))
ORDEM = {'Dúvida': 0, 'Compara': 1, 'Não compara': 2, 'Só 2026': 3, 'Só 2027': 4}
linhas.sort(key=lambda l: (ORDEM[l['leitura']], -max(l['v26'], l['v27'])))

F = 'Arial'
AZUL, BRANCO = '151E49', 'FFFFFF'
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
fino = Side(style='thin', color='E2E8F0')
borda = Border(bottom=fino)
MOEDA = '"R$" #,##0;-"R$" #,##0;"–"'

wb = Workbook()
ws = wb.active
ws.title = 'Comparação'
ws['A1'] = 'Manutenção de imóveis — Mega Curitiba · itens do Orç 2026 × Orç 2027'
ws['A1'].font = Font(name=F, bold=True, size=13, color=AZUL)
ws['A2'] = ('Preencha a coluna COMPARA? com SIM ou NÃO nas linhas amarelas (dúvidas). As outras já vêm com a leitura da '
            'análise e podem ser trocadas. Só os pares com SIM entram na comparação item a item da apresentação.')
ws['A2'].font = Font(name=F, size=9, color='475569')
ws['A3'] = ('Orç 2026: "Modelos 2025 Megas" (centro de custo do condomínio) + "2025 - Contratos" (valores de 2026). '
            'Orç 2027: modelo 090 + planilha de contratos de manutenção. Totais iguais aos da METRAGEM-COND.')
ws['A3'].font = Font(name=F, size=9, color='94A3B8')

cab = ['#', 'Tipo', 'Item no Orç 2026', 'Orç 2026', 'Item(ns) no Orç 2027', 'Categoria 2027', 'Orç 2027', 'Δ R$',
       'Leitura da análise', 'Por quê', 'COMPARA?', 'Comentário do gestor']
H = 5
for c, t in enumerate(cab, 1):
    cel = ws.cell(row=H, column=c, value=t)
    cel.font = Font(name=F, bold=True, size=9, color=BRANCO)
    cel.fill = fill('F97316' if t == 'COMPARA?' else AZUL)
    cel.alignment = Alignment(vertical='center', horizontal='center' if c in (1, 4, 7, 8, 11) else 'left', wrap_text=True)
ws.row_dimensions[H].height = 26

PREENCHE = {'Compara': 'SIM', 'Não compara': 'NÃO'}
COR = {'Dúvida': 'FEF3C7', 'Só 2026': 'F8FAFC', 'Só 2027': 'F8FAFC'}
for i, l in enumerate(linhas):
    r = H + 1 + i
    vals = [i + 1, l['tipo'], l['d26'] or '—', l['v26'] or None, l['d27'] or '—', l['cat27'] or '', l['v27'] or None,
            '=G{0}-D{0}'.format(r), l['leitura'], l['motivo'], PREENCHE.get(l['leitura']), None]
    for c, v in enumerate(vals, 1):
        cel = ws.cell(row=r, column=c, value=v)
        cel.font = Font(name=F, size=9, bold=(c == 11), color='94A3B8' if l['leitura'] in ('Só 2026', 'Só 2027') and c not in (4, 7) else '151E49')
        cel.alignment = Alignment(vertical='top', wrap_text=c in (3, 5, 10, 12), horizontal='center' if c in (1, 11) else None)
        cel.border = borda
        if c in (4, 7, 8): cel.number_format = MOEDA
        if l['leitura'] in COR: cel.fill = fill(COR[l['leitura']])
    if l['leitura'] == 'Dúvida':
        ws.cell(row=r, column=11).fill = fill('FDE68A')
ult = H + len(linhas)

dv = DataValidation(type='list', formula1='"SIM,NÃO"', allow_blank=True, showDropDown=False)
dv.error = 'Use SIM ou NÃO'
ws.add_data_validation(dv)
dv.add('K{0}:K{1}'.format(H + 1, ult + 200))

for col, w in zip('ABCDEFGHIJKL', [5, 10, 42, 13, 52, 18, 13, 13, 14, 60, 12, 30]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'D6'
ws.auto_filter.ref = 'A{0}:L{1}'.format(H, ult)

# ---- Resumo: tudo por fórmula sobre a aba de comparação ----
rs = wb.create_sheet('Resumo')
rs['A1'] = 'Resumo da comparação (atualiza sozinho com as marcações)'
rs['A1'].font = Font(name=F, bold=True, size=12, color=AZUL)
C = "'Comparação'!"
itens = [
    ('Total Orç 2026 (R$)', '=SUM({0}D6:D1000)'.format(C)),
    ('Total Orç 2027 (R$)', '=SUM({0}G6:G1000)'.format(C)),
    ('Itens marcados SIM — Orç 2026 (R$)', '=SUMIFS({0}D6:D1000,{0}K6:K1000,"SIM")'.format(C)),
    ('Itens marcados SIM — Orç 2027 (R$)', '=SUMIFS({0}G6:G1000,{0}K6:K1000,"SIM")'.format(C)),
    ('Variação dos itens que comparam (R$)', '=B6-B5'),
    ('Variação dos itens que comparam (%)', '=IF(B5=0,0,B6/B5-1)'),
    ('Pares marcados SIM', '=COUNTIF({0}K6:K1000,"SIM")'.format(C)),
    ('Dúvidas ainda sem resposta', '=COUNTIFS({0}I6:I1000,"Dúvida",{0}K6:K1000,"")'.format(C)),
]
for i, (rot, f) in enumerate(itens):
    r = 3 + i
    a = rs.cell(row=r, column=1, value=rot); a.font = Font(name=F, size=10)
    b = rs.cell(row=r, column=2, value=f); b.font = Font(name=F, size=10, bold=True, color=AZUL)
    b.number_format = '0.0%' if '(%)' in rot else ('#,##0' if 'Pares' in rot or 'Dúvidas' in rot else MOEDA)
rs['A12'] = 'Fonte: aba Comparação (Orç 2026 = Modelos 2025 Megas + 2025 - Contratos; Orç 2027 = modelo 090 e contratos de Curitiba).'
rs['A12'].font = Font(name=F, size=8, color='94A3B8')
rs.column_dimensions['A'].width = 40
rs.column_dimensions['B'].width = 18

os.makedirs(os.path.join(S, 'saida'), exist_ok=True)
out = os.path.join(S, 'saida', 'comparacao_itens.xlsx')
wb.save(out)
print(out, os.path.getsize(out))
