"""Incorpora a entrevista com o financeiro (08/10/2026) na planilha de diagnóstico.

Roda uma vez sobre a planilha da Fase 1 (4 entrevistas). Mantém a formatação existente
copiando o estilo de linhas vizinhas para as linhas novas.
Uso: python3 atualizar_planilha_financeiro.py <arquivo.xlsx>
"""
import sys
from copy import copy
import openpyxl

ARQ = sys.argv[1]
wb = openpyxl.load_workbook(ARQ)


def copiar_estilo(ws, origem, destino, ncols):
    for c in range(1, ncols + 1):
        a, b = ws.cell(origem, c), ws.cell(destino, c)
        b.font, b.fill, b.border = copy(a.font), copy(a.fill), copy(a.border)
        b.alignment, b.number_format = copy(a.alignment), a.number_format
    ws.row_dimensions[destino].height = ws.row_dimensions[origem].height


def linha(ws, rotulo, col=1):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == rotulo:
            return r
    raise KeyError(rotulo)


# ---------------------------------------------------------------- Visão Geral
ws = wb['Visão Geral']
ws['A2'] = 'Capital Realty · Diagnóstico (Fase 1) · Documento vivo · atualizado em 08/10/2026'
ws.cell(linha(ws, 'Fase atual'), 2).value = 'Fase 1 — Diagnóstico (entrevistas concluídas; falta consolidar a ata)'
ws.cell(linha(ws, 'Entrevistas concluídas'), 2).value = '5: colaboradores Wilson, Cadu, Jonatas, Ricardo + gerente do financeiro (08/10/2026)'
ws.cell(linha(ws, 'Pendências'), 2).value = (
    'Confirmar 6 pontos com o financeiro (nome, regra do dia 25, documento do Thiago, "ticket", '
    'quem mais consultar, valores vigentes); obter o PDF do procedimento 2022; consolidar a ata de diagnóstico.')

# --------------------------------------------------------- Dores Consolidadas
ws = wb['Dores Consolidadas']
ws['A1'] = 'Painel de Dores — Consolidado (4 colaboradores + financeiro)'
ws['A2'] = 'Prioridade baseada em nº de confirmações e impacto · financeiro entrevistado em 08/10/2026'

def dor(n):
    return linha(ws, n)

atual = {
    1: (None, '4/4 + Fin.', 'Todos + Financeiro',
        "Regras 'por item' e não 'por conceito': chocolate não, sonho sim; 'cueca virada' é doce?; água só em certo horário. "
        "FINANCEIRO: item não reembolsável é o 2º erro mais comum; a regra de refeição foi redefinida pelo Thiago e não está no procedimento; "
        "para o financeiro sonho NÃO pode (Jonatas entendia que podia) — prova de que a regra conhecida não é a aplicada."),
    2: ('Consulta informal — de mão dupla', '4/4 + Fin.', 'Todos + Financeiro',
        'Colaboradores consultam antes de lançar (Ricardo escalou à diretoria). E a MAIOR DOR DO FINANCEIRO é o inverso: '
        'perguntar ao colaborador o que está na NF.'),
    5: (None, 'Você + 4/4 + Fin.', 'Todos + Financeiro',
        'Documento de 2022. FINANCEIRO: KM reajustado todo ano por e-mail (gasolina + IPVA); limites de passagem/hotel revistos todo ano; '
        'refeição revisada pelo Thiago; CNPJ não escrito. As regras vivas estão fora do documento. Wilson nunca teve acesso; Cadu e Ricardo não acham.'),
    6: (None, '2/4 + Fin.', 'Cadu, Ricardo (+ Jonatas) + Financeiro',
        'Cupom que não vale como NF, ou sistema do fornecedor que falha. FINANCEIRO: CNPJ obrigatório para estacionamento e combustível; '
        'jantar é exceção. Regra esclarecida, falta escrever.'),
    7: (None, '2/4 + Fin.', 'Cadu, Ricardo (reforço Jonatas) + Financeiro',
        'REENQUADRADA: o prazo vem da competência contábil e é INEGOCIÁVEL (financeiro). A melhoria não é afrouxar o prazo, '
        'é ajudar a cumpri-lo — ver dor #12.'),
    9: (None, '2/4 + Fin.', 'Jonatas, Ricardo + Financeiro',
        'KM calculado em planilha externa ao Fluig. FINANCEIRO: a data do formulário de KM diverge da data no corpo do Fluig — erro recorrente.'),
    10: (None, '1/4 + Fin.', 'Jonatas + Financeiro',
         'Reembolsos em lote: uma linha errada bloqueia todo o processo. FINANCEIRO confirmou: não dá para aprovar parcialmente.'),
    11: ('Controle gerencial: sem controle de orçamento', 'Confirmada em parte', 'Financeiro',
         'Status visível no Fluig: sim. Controle de orçamento de viagens por área/colaborador: NÃO EXISTE. '
         'Volume (30–40/mês) e taxa de erro (3 em 10) só de cabeça.'),
}
for n, (titulo, conf, fontes, det) in atual.items():
    r = dor(n)
    if titulo: ws.cell(r, 2).value = titulo
    ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value = conf, fontes, det

r11 = dor(11)
ws.cell(r11, 3).value = 'BAIXA'

modelo = {'ALTA': dor(1), 'MÉDIA': dor(6), 'BAIXA': r11}
novas = [
    (12, 'Lançamento fora da competência', 'ALTA', 'Fin.: erro mais comum', 'Financeiro',
     'Registro por competência (inegociável). Corte no dia 25; a OC é gerada no dia em que o processo termina. Quem abre no dia 26, '
     'ou leva reprovações e ajustes, cai no mês seguinte e perde a OC. Ex.: viagem 02/09, reembolso aberto 26/09.'),
    (13, 'Descrição não diz o que foi comprado', 'MÉDIA', 'Fin.: maior dor', 'Financeiro',
     'O financeiro precisa perguntar o que está na NF. A data da viagem também costuma faltar.'),
    (14, 'Retrabalho alto (~3 em 10 com erro)', 'MÉDIA', 'Financeiro', 'Financeiro',
     '30–40 solicitações por mês, ~3 de cada 10 chegam com erro (≈10/mês). Cada reprovação consome dias do prazo e alimenta a dor #12.'),
    (15, 'O processo cobre mais que viagens', 'BAIXA', 'Financeiro', 'Financeiro',
     'O mesmo fluxo é usado para despesas que não são de viagem (ex.: eventos do marketing); o procedimento só trata de viagens.'),
]
r = ws.max_row
for n, titulo, prio, conf, fontes, det in novas:
    r += 1
    copiar_estilo(ws, modelo[prio], r, 6)
    for c, v in enumerate([n, titulo, prio, conf, fontes, det], 1):
        ws.cell(r, c).value = v
# BAIXA usa a cor cinza do antigo "A VALIDAR"; a #12 é ALTA e copia o vermelho.

# --------------------------------------------------------- Matriz Dor x Pessoa
ws = wb['Matriz Dor x Pessoa']
for rng in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(rng))
# Total vai de F para G; F vira Financeiro.
for row in range(3, ws.max_row + 1):
    f, g = ws.cell(row, 6), ws.cell(row, 7)
    g.value = f.value
    g.font, g.fill, g.border, g.alignment = copy(f.font), copy(f.fill), copy(f.border), copy(f.alignment)
    e = ws.cell(row, 5)
    f.font, f.fill, f.border, f.alignment = copy(e.font), copy(e.fill), copy(e.border), copy(e.alignment)
    f.value = None
ws.column_dimensions['F'].width = 14
ws.column_dimensions['G'].width = 18
ws.merge_cells('A1:G1'); ws.merge_cells('A2:G2')
ws['F3'] = 'Financeiro'
fin = {
    'Lista do que é reembolsável': ('✔', '5/5'),
    'Consulta informal antes de lançar': ('✔ (inverso)', '5/5'),
    'Falta de mobilidade (só PC/VPN)': ('—', None),
    'Desistiu de solicitar por burocracia': ('—', None),
    'Procedimento desatualizado / difícil acesso': ('✔', '5/5'),
    'Documento fiscal inválido (cupom/NF/CNPJ)': ('✔', '3/5 +1'),
    'Prazo rígido': ('✔', '3/5 +1'),
    'Digitalização manual de notas': ('—', '2/5'),
    'Quilometragem (planilha externa/erros)': ('✔', '3/5'),
    'Erro em 1 linha trava o pacote': ('✔', '2/5'),
    'Desistiu de solicitar?': ('—', None),
    'Prefere manter Fluig?': ('Sim', '4x Fluig / 1 app'),
}
for row in range(4, ws.max_row + 1):
    k = ws.cell(row, 1).value
    if k in fin:
        v, tot = fin[k]
        ws.cell(row, 6).value = v
        if tot: ws.cell(row, 7).value = tot
ws.cell(linha(ws, 'Prazo rígido'), 1).value = 'Prazo rígido / fora da competência'
# Linhas novas acima das duas perguntas de perfil (desistiu / Fluig).
pos = linha(ws, 'Desistiu de solicitar?')
ws.insert_rows(pos, 2)
for i, (nome, vals) in enumerate([
        ('Descrição não diz o que foi comprado', ['—', '—', '—', '—', '✔', '1/5']),
        ('Controle de orçamento', ['—', '—', '—', '—', '✔ (não existe)', '—'])]):
    rr = pos + i
    copiar_estilo(ws, pos - 1, rr, 7)
    ws.cell(rr, 1).value = nome
    for c, v in enumerate(vals, 2):
        ws.cell(rr, c).value = v

# ------------------------------------------------------------- aba Financeiro
src = wb['Ricardo']
ws = wb.copy_worksheet(src)
ws.title = 'Financeiro'
wb.move_sheet(ws, offset=wb.sheetnames.index('Ricardo') + 1 - wb.sheetnames.index('Financeiro'))
ws.sheet_properties.tabColor = src.sheet_properties.tabColor
qa = [
    ('1.1 Processo e atores', 'Os colaboradores que viajam.'),
    ('1.2 Volume mensal', '30 a 40 solicitações por mês.'),
    ('1.3 Etapas fora do Fluig', 'KM é feito por planilha.'),
    ('2.1 Erros mais comuns', 'Fora da competência (o mais comum — ex.: solicitar em setembro despesa de julho); itens não reembolsáveis; '
                              'não descrever a data da viagem; data do formulário de KM diferente da do corpo do Fluig.'),
    ('2.2 Incompletas / fora do prazo', 'De cada 10, umas 3 chegam erradas.'),
    ('2.3 Despesas que geram dúvida', 'Café da manhã: bebidas lácteas e proteína sim, sonhos não. Não pode ser exclusivamente doce; no jantar salgado com doce pode. '
                                      'Na viagem, parar para um salgado e seguir é aceito, mas priorizar alimentos. O Thiago definiu o que é refeição.'),
    ('2.4 Maior dor do financeiro', 'Perguntar o que foi descrito na NF — não fica claro o que está sendo solicitado.'),
    ('2.5 Linha errada / aprovação parcial', 'Não dá para aprovar parcialmente.'),
    ('3.1 R$ 0,97/km', 'Reajustado anualmente (gasolina + IPVA), comunicado por e-mail.'),
    ('3.2 Limites passagem/hotel', 'Revistos anualmente; deveriam ser seguidos.'),
    ('3.3 Regra de refeições', 'Existe; o Thiago revisou.'),
    ('3.4 Adiantamento R$ 200', 'Vigente.'),
    ('3.5 CNPJ na nota', 'Estacionamento e combustível: sim. Deveria ser regra, mas jantar é exceção.'),
    ('4.1 Status / painel', 'Sim, pelo Fluig.'),
    ('4.2 Tempo solicitação → pagamento', '10 dias.'),
    ('4.3 Controle de orçamento', 'Não controla.'),
    ('5.1 O que mudaria', 'Incluir no Fluig as informações de suporte; painel com FAQ mostrando valores e a regra da data.'),
    ('5.2 Ferramentas / preferência', 'Nenhuma avaliada. Preferência pelo Fluig, sem restrição.'),
    ('5.3 Restrições do redesenho', 'Fora da competência é inegociável; adiantamento sem NF não é negociável. Norma contábil: registro por competência, não por caixa.'),
    ('5.4 Quem mais consultar', '(sem resposta)'),
    ('Outras observações', 'Corte adaptado para o dia 25; se viajar depois, lançar no primeiro dia. A OC é gerada no dia em que o processo termina: '
                           'quem solicita dia 26, com reprovações e ajustes, fica fora e perde a OC (ex.: viagem 02/09, reembolso aberto 26/09). '
                           'Processo também usado para despesas que não são de viagem (eventos do marketing). Ideias: painel fixo/FAQ, '
                           '"visualizador", lembrete de viagens. Anotado sem detalhe: "ticket", "pessoas assumindo o risco".'),
]
ws['A1'] = 'Entrevista — Financeiro'
ws['A2'] = 'Gerente do financeiro (nome a registrar) · 08/10/2026 · ~33 min · roteiro do financeiro'
base = 4  # linha 3 é o cabeçalho Pergunta/Resposta
modelo_r = 5
for i, (p, resp) in enumerate(qa):
    rr = base + i
    if rr > 17:
        copiar_estilo(ws, modelo_r, rr, 2)
    ws.cell(rr, 1).value, ws.cell(rr, 2).value = p, resp
    ws.row_dimensions[rr].height = None  # deixa o Excel ajustar ao texto

# -------------------------------------------------------------------- Preservar
ws = wb['Preservar']
r = linha(ws, 'Prazo de pagamento')
ws.cell(r, 2).value = 'Após aprovação, atende os quatro (3–5 dias). Elogio recorrente. Financeiro mede 10 dias desde a solicitação (inclui a aprovação).'
ws.cell(r, 3).value = 'Todos + Financeiro'
r = linha(ws, 'Visibilidade no Fluig')
ws.cell(r, 2).value = 'Solicitante e financeiro conseguem ver o status da solicitação.'
ws.cell(r, 3).value = 'Todos + Financeiro'
r = linha(ws, 'Plataforma Fluig')
ws.cell(r, 2).value = 'DECIDIDO: 4 de 5 preferem o Fluig; o financeiro não tem restrição e nenhuma outra ferramenta foi avaliada. Direção: evoluir o Fluig.'
ws.cell(r, 3).value = 'Wilson, Cadu, Jonatas, Financeiro x Ricardo'

# ------------------------------------------------------------ Ideias de Solução
ws = wb['Ideias de Solução']
r = linha(ws, 'Adiantamento / verba fixa por viagem')
ws.cell(r, 2).value = ws.cell(r, 2).value + ' ⚠️ Esbarra na regra do financeiro: adiantamento sem NF não é negociável.'
r = linha(ws, 'KM integrado (sem planilha externa)')
ws.cell(r, 3).value = 'Jonatas, Ricardo, Financeiro'
novas_ideias = [
    ('FAQ / painel de regras no Fluig', 'Painel fixo com valores vigentes, regra da data, pode/não pode e dúvidas frequentes — um "visualizador" das regras.', 'Financeiro'),
    ('Informações de suporte no formulário', 'Explicação e exemplos dentro do próprio formulário do Fluig.', 'Financeiro'),
    ('Lembrete de viagem / de prazo', 'Aviso para lançar antes do dia 25 e não perder a competência.', 'Financeiro'),
    ('Aprovação parcial', 'Aprovar as linhas certas e devolver só a errada, em vez de travar o pacote.', 'Jonatas, Financeiro'),
]
ult = ws.max_row
for i, vals in enumerate(novas_ideias, 1):
    copiar_estilo(ws, ult, ult + i, 3)
    for c, v in enumerate(vals, 1):
        ws.cell(ult + i, c).value = v

# -------------------------------------------------------------- Próximos Passos
ws = wb['Próximos Passos']
r = linha(ws, 'Entrevista com o time financeiro')
ws.cell(r, 2).value = 'Feita em 08/10/2026 (ver aba Financeiro e Ata_Entrevista_Financeiro.md). Falta confirmar 6 pontos: nome da entrevistada, regra exata do dia 25, documento do Thiago sobre refeição, "ticket" e "pessoas assumindo o risco", quem mais consultar, e-mails com valores vigentes.'
ws.cell(r, 3).value = 'Feita · confirmar'
r = linha(ws, 'Definir direção de plataforma')
ws.cell(r, 2).value = 'Decidido: evoluir o Fluig (4 de 5 + financeiro sem restrição). Falta levantar com o TI o que o Fluig permite (mobilidade sem VPN, FAQ, lembrete, aprovação parcial, KM integrado).'
ws.cell(r, 3).value = 'Decidido'
r = linha(ws, 'Redesenhar regras de itens elegíveis')
ws.cell(r, 2).value = "Migrar de 'por item' para 'por conceito' partindo da definição do Thiago; tabela pode/não pode; regra de competência e do dia 25 com exemplo; incluir despesas que não são de viagem."
ult = ws.max_row
copiar_estilo(ws, ult, ult + 1, 3)
ws.cell(ult + 1, 1).value = 'Obter o PDF do procedimento 2022'
ws.cell(ult + 1, 2).value = 'Documento base não foi achado no Drive. Necessário para reescrever o procedimento (Fase 3/4).'
ws.cell(ult + 1, 3).value = 'Pendente'

wb.save(ARQ)
print('ok', wb.sheetnames)
